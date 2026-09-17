#!/usr/bin/env python3
"""Extract the teacher's own diagrams out of the uploaded material, as images.

Why this exists
---------------
`tools/audit_dcc_coverage.py` measures the gap: 175 diagram-bearing slides and 65
diagram-bearing PDF pages against the portal's 20 hand-built SVG figures. The
figures that exist are clean and themable, but they are my redrawing of a handful
of concepts; everything else the teacher drew was simply not on the site.

This tool pulls those drawings out as images so they can be shown as they were
taught. The two formats need different treatment, and the difference matters:

* **Slide decks.** A picture on a slide is a picture shape with an image blob
  behind it, so the blob is taken whole. Groups are recursed into because a
  diagram is often a group of fragments.

* **PDFs.** The embedded image objects are NOT whole diagrams. Measured on
  `REST_SOAP_webservices_Lecture.pdf`: 204 image objects across 15 pages, about
  fourteen per page, because a diagram exported from a drawing tool arrives as
  sliced bitmap fragments. Extracting those objects would produce a pile of
  unreadable slices, so a PDF's diagram page is RENDERED instead - whole page, at
  a resolution where its smallest label is still legible.

Filtering is the other half of the job: a deck's picture shapes include logos,
bullet glyphs and decorative rules, and those are not diagrams.

    python tools/extract_dcc_figures.py --report     # what would be kept, and why
    python tools/extract_dcc_figures.py              # write the images + manifest

Output: assets/dcc-slides/ch<N>/<slug>.webp, plus FIGURES.json describing every
image (unit, source file, slide or page, caption, alt text).
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DCC = ROOT / 'dcc'
OUT = ROOT / 'assets' / 'dcc-slides'
MANIFEST = OUT / 'FIGURES.json'

# --- what counts as a diagram -------------------------------------------------
# A bullet glyph, a logo and a hairline rule are all pictures on a slide. These
# floors are what separates them from a drawing: measured across the decks, the
# decks' decorative art sits well under all four.
MIN_BLOB = 6_000            # bytes of source image
MIN_PX_W = 300              # pixels, the image's own width
MIN_PX_H = 130              # pixels, the image's own height
MIN_EMU_W = 2_286_000       # 2.5 inches displayed on the slide (914400 EMU per inch)
MIN_AREA_FRAC = 0.02        # or at least this share of the slide's area

# Rendering a PDF page: 150 DPI puts a 720pt-wide slide at about 1500px, where a
# 10pt label is still ~21px tall. Below ~110 DPI the OCR-sized text in these
# decks starts to blur.
PDF_DPI = 150
WEBP_QUALITY = 78
MAX_W = 1600                # nothing gains from being wider than the reading column twice over


def slug(text: str) -> str:
    """A filename fragment. Lowercase FIRST, then substitute separators - doing it
    the other way round turns every capital into a separator."""
    text = text.lower()
    text = re.sub(r'[^a-z0-9]+', '-', text)
    return text.strip('-')[:60] or 'image'


def chapter_of(path: Path) -> int | None:
    for part in path.parts:
        m = re.match(r'chapter[ _]?(\d+)', part, re.I)
        if m:
            return int(m.group(1))
    return None


def is_diagram(w: int, h: int, blob_len: int, emu_w: int | None, slide_area: int | None) -> tuple[bool, str]:
    if blob_len < MIN_BLOB:
        return False, f'blob {blob_len}B < {MIN_BLOB}'
    if w < MIN_PX_W:
        return False, f'{w}px wide < {MIN_PX_W}'
    if h < MIN_PX_H:
        return False, f'{h}px tall < {MIN_PX_H}'
    if emu_w is not None and emu_w < MIN_EMU_W:
        return False, f'displayed {emu_w / 914400:.1f}in < 2.5in'
    if emu_w is None and slide_area:
        if (w * h) / slide_area < MIN_AREA_FRAC:
            return False, f'{100 * (w * h) / slide_area:.1f}% of slide area'
    return True, 'kept'


def encode(img, src_w: int, src_h: int) -> tuple[bytes, int, int]:
    """Re-encode to WebP, returning the bytes AND the final pixel size.

    The size is returned rather than recomputed because the manifest carries it
    into the page as width/height attributes, which is what stops a column of
    lazy images from jumping as each one arrives. An image wider than MAX_W is
    resized, so its source size is not its final size - recording the source
    would put a wrong height in the markup for exactly the widest pictures.
    """
    from PIL import Image
    if img.mode not in ('RGB', 'RGBA'):
        img = img.convert('RGB')
    if img.width > MAX_W:
        img = img.resize((MAX_W, round(img.height * MAX_W / img.width)), Image.LANCZOS)
    buf = io.BytesIO()
    # WebP keeps a screenshot's flat fills small and its text edges sharp at a
    # fraction of a PNG, which matters at a few hundred images.
    img.save(buf, format='WEBP', quality=WEBP_QUALITY, method=5)
    return buf.getvalue(), img.width, img.height


def picture_shapes(shapes, depth: int = 0):
    """Every shape carrying an image blob, groups recursed into.

    A `Picture` (shape_type 13) is the obvious case, and testing for it alone is
    what this tool did first - which silently dropped the *placeholder* pictures.
    In `Ref_CloudComptng.pptx`, 28 of that deck's 31 diagrams are
    `PlaceholderPicture` (shape_type 14) rather than `Picture`, so the deck
    yielded three images where the audit counted thirty-one. A deck built from a
    template put its drawings in the template's picture placeholders, and both
    classes expose `.image`; that is the thing worth testing.
    """
    for shape in shapes:
        try:
            if shape.shape_type == MSO_SHAPE_TYPE_GROUP and depth < 3:
                yield from picture_shapes(shape.shapes, depth + 1)
                continue
        except Exception:
            pass
        try:
            if shape.image is not None:
                yield shape
        except (AttributeError, ValueError, KeyError):
            continue  # not a picture, or a picture with no blob behind it


# Imported lazily so the module can be read without python-pptx installed.
try:
    from pptx.enum.shapes import MSO_SHAPE_TYPE as _MSO
    MSO_SHAPE_TYPE_GROUP = _MSO.GROUP
except Exception:                                     # pragma: no cover
    MSO_SHAPE_TYPE_GROUP = 6


def deck_files() -> list[tuple[Path, Path]]:
    """(path to read, path that names the unit) for every deck.

    The second element differs from the first for a legacy .ppt: it is converted
    through PowerPoint into a cache directory, which no longer knows which unit
    its source belonged to, so the original path is carried alongside.
    """
    pairs: list[tuple[Path, Path]] = []
    for path in sorted(DCC.rglob('*.pptx')):
        pairs.append((path, path))
    legacy = sorted(DCC.rglob('*.ppt'))
    if legacy:
        sys.path.insert(0, str(ROOT / 'tools'))
        from dcc_extract import convert_ppt
        for path in legacy:
            pairs.append((convert_ppt(path), path))
    return pairs


def best_caption(shapes) -> str:
    """The most descriptive text on the slide, for the image's caption.

    Taking the first text frame looked right and measured badly: 67 of 231
    drawings got a usable caption from it and the rest got a section heading
    ("1.1 Introduction"), a page number, or "Contd..". A slide's title
    placeholder is the intent when it has one, and otherwise the most words
    carry the most meaning - so the longest frame in a sane range wins.
    """
    frames: list[tuple[bool, str]] = []
    for shape in shapes:
        if not shape.has_text_frame:
            continue
        text = shape.text_frame.text.strip().replace('\n', ' ')
        if not text:
            continue
        is_title = False
        try:
            is_title = bool(shape.is_placeholder and shape.placeholder_format.idx == 0)
        except Exception:
            pass
        frames.append((is_title, text))
    if not frames:
        return ''
    for is_title, text in frames:
        if is_title and len(text) >= 6:
            return text[:180]
    # "Contd.." and friends say nothing the slide number does not already say.
    filler = re.compile(r'^(cont(inued)?\.*|slide \d+|\d+)$', re.I)
    useful = [t for _, t in frames if 15 <= len(t) <= 180 and not filler.match(t)]
    if useful:
        return max(useful, key=len)[:180]
    return frames[0][1][:180]


def from_decks(report: bool) -> list[dict]:
    from PIL import Image
    from pptx import Presentation

    found: list[dict] = []
    seen_hashes: set[str] = set()
    for read_path, origin in deck_files():
        deck = origin
        unit = chapter_of(deck)
        if unit is None:
            continue
        prs = Presentation(str(read_path))
        slide_area = (prs.slide_width or 0) * (prs.slide_height or 0)
        for idx, slide in enumerate(prs.slides, 1):
            # The slide's own text is the best caption/alt text available: it is
            # what the teacher wrote next to the drawing.
            texts = [sh.text_frame.text.strip() for sh in slide.shapes
                     if sh.has_text_frame and sh.text_frame.text.strip()]
            caption = best_caption(slide.shapes)
            body = ' '.join(texts)[:600]

            for shape in picture_shapes(slide.shapes):
                try:
                    blob = shape.image.blob
                    img = Image.open(io.BytesIO(blob))
                    w, h = img.size
                except Exception:
                    continue
                emu = getattr(shape, 'width', None)
                keep, why = is_diagram(w, h, len(blob), emu, slide_area)
                digest = hashlib.sha256(blob).hexdigest()
                # The same logo or bullet appears on every slide; keep one.
                if digest in seen_hashes:
                    keep, why = False, 'duplicate of an earlier image'
                if not keep:
                    # Recorded in BOTH modes. Skipping it outside --report made
                    # the write run summarise itself as `rejected 0`, which
                    # reads as "nothing was filtered" when 63 things were.
                    found.append({'unit': unit, 'kept': False, 'why': why,
                                  'source': deck.name, 'slide': idx,
                                  'px': f'{w}x{h}', 'bytes': len(blob)})
                    continue
                seen_hashes.add(digest)
                data, ow, oh = encode(img, w, h) if not report else (b'', w, h)
                found.append({
                    'unit': unit, 'kept': True, 'why': why, 'kind': 'slide',
                    'source': deck.name, 'slide': idx, 'px': f'{ow}x{oh}',
                    'bytes': len(data), 'src_bytes': len(blob),
                    'caption': caption, 'body': body,
                    'name': f'ch{unit}/{slug(deck.stem)}-s{idx:02d}-{len(seen_hashes):03d}.webp',
                    '_data': data,
                })
    return found


def from_pdfs(report: bool) -> list[dict]:
    import pypdfium2 as pdfium
    from PIL import Image

    PATH_OBJ = pdfium.raw.FPDF_PAGEOBJ_PATH
    found: list[dict] = []
    for pdf in sorted(DCC.rglob('*.pdf')):
        if 'Books' in str(pdf) or 'syllabus' in pdf.name.lower():
            continue
        unit = chapter_of(pdf)
        if unit is None:
            continue
        doc = pdfium.PdfDocument(str(pdf))
        for i in range(len(doc)):
            page = doc[i]
            # The audit's signal: a page dense in vector paths is a drawing.
            if len(list(page.get_objects(filter=(PATH_OBJ,)))) <= 40:
                continue
            bitmap = page.render(scale=PDF_DPI / 72)
            img = bitmap.to_pil()
            data, ow, oh = (b'', img.width, img.height) if report else encode(img, img.width, img.height)
            # The page's own text is the caption: on a rendered page it is part
            # of the picture, so a reader needs it read out separately.
            text = (page.get_textpage().get_text_range() or '').strip()
            first = next((ln.strip() for ln in text.splitlines() if ln.strip()), '')
            found.append({
                'unit': unit, 'kept': True, 'why': 'diagram page', 'kind': 'pdf',
                'source': pdf.name, 'page': i + 1, 'px': f'{ow}x{oh}',
                'bytes': len(data), 'caption': first[:180], 'body': text[:600],
                'name': f'ch{unit}/{slug(pdf.stem)}-p{i + 1:02d}.webp',
                '_data': data,
            })
        doc.close()
    return found


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--report', action='store_true',
                    help='list what would be kept and rejected, writing nothing')
    args = ap.parse_args()

    slides = from_decks(args.report)
    pdfs = from_pdfs(args.report)
    everything = slides + pdfs
    kept = [f for f in everything if f['kept']]

    print(f'{"candidates":<24}{len(everything):>6}')
    print(f'{"kept":<24}{len(kept):>6}')
    print(f'{"rejected":<24}{len(everything) - len(kept):>6}')
    print()
    print('rejected, by reason:')
    reasons: dict[str, int] = {}
    for f in everything:
        if not f['kept']:
            r = re.sub(r'[\d.]+', 'N', f['why'])
            reasons[r] = reasons.get(r, 0) + 1
    for r, n in sorted(reasons.items(), key=lambda kv: -kv[1]):
        print(f'  {r:<40} {n:>5}')
    print()
    print(f'{"unit":<8}{"slides":>8}{"pdf pages":>11}{"images":>8}{"webp MB":>10}')
    total_bytes = 0
    for unit in range(1, 10):
        group = [f for f in kept if f['unit'] == unit]
        s = sum(1 for f in group if f['kind'] == 'slide')
        p = sum(1 for f in group if f['kind'] == 'pdf')
        b = sum(f['bytes'] for f in group)
        total_bytes += b
        print(f'  ch{unit:<5}{s:>8}{p:>11}{len(group):>8}{b / 1e6:>10.2f}')
    print(f'  {"TOTAL":<6}{sum(1 for f in kept if f["kind"] == "slide"):>8}'
          f'{sum(1 for f in kept if f["kind"] == "pdf"):>11}{len(kept):>8}'
          f'{total_bytes / 1e6:>10.2f}')

    if args.report:
        print()
        print('samples:')
        for f in kept[:12]:
            print(f"  ch{f['unit']}  {f['source'][:34]:<36} "
                  f"{('slide ' + str(f['slide'])) if f['kind'] == 'slide' else ('page ' + str(f['page'])):<10}"
                  f" {f['px']:>10}  {f['caption'][:40]}")
        return 0

    # --- write ---
    if OUT.exists():
        for old in OUT.rglob('*.webp'):
            old.unlink()
    for f in kept:
        dest = OUT / f['name']
        dest.parent.mkdir(parents=True, exist_ok=True)
        dest.write_bytes(f['_data'])
    manifest = [{k: v for k, v in f.items() if not k.startswith('_')} for f in kept]
    OUT.mkdir(parents=True, exist_ok=True)
    MANIFEST.write_text(json.dumps(manifest, indent=1, ensure_ascii=False) + '\n',
                        encoding='utf-8')
    print()
    print(f'wrote {len(kept)} image(s) into {OUT.relative_to(ROOT)} and {MANIFEST.name}')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
